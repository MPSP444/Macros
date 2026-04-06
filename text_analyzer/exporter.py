from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os


def _register_fonts():
    """Try to register a Unicode-capable font for Russian text."""
    font_candidates = [
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/truetype/freefont/FreeSans.ttf',
        '/usr/share/fonts/TTF/DejaVuSans.ttf',
        '/System/Library/Fonts/Supplemental/Arial Unicode.ttf',
        'C:/Windows/Fonts/arial.ttf',
    ]
    for path in font_candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont('UniFont', path))
                return 'UniFont'
            except Exception:
                continue
    return 'Helvetica'


class ReportExporter:
    def export_pdf(self, filepath: str, data: dict):
        font_name = _register_fonts()

        doc = SimpleDocTemplate(
            filepath,
            pagesize=A4,
            rightMargin=2 * cm,
            leftMargin=2 * cm,
            topMargin=2 * cm,
            bottomMargin=2 * cm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontName=font_name,
            fontSize=16,
            spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontName=font_name,
            fontSize=12,
            spaceAfter=8,
        )
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontName=font_name,
            fontSize=10,
            spaceAfter=4,
        )

        story = []
        story.append(Paragraph('TextAnalyzer Pro — Отчёт', title_style))
        story.append(Spacer(1, 0.5 * cm))

        # Score table
        story.append(Paragraph('Метрики анализа', heading_style))
        score = data.get('score', {})
        score_data = [
            ['Метрика', 'Значение'],
            ['Общий Score', str(score.get('total', '—'))],
            ['TTR', str(score.get('ttr', '—'))],
            ['Burstiness', str(score.get('burstiness', '—'))],
            ['Средняя длина предложения', str(score.get('avg_sentence_length', '—'))],
            ['Кол-во слов', str(score.get('word_count', '—'))],
            ['Уникальных слов', str(score.get('unique_words', '—'))],
            ['Кол-во предложений', str(score.get('sentence_count', '—'))],
        ]
        score_table = Table(score_data, colWidths=[10 * cm, 6 * cm])
        score_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(score_table)
        story.append(Spacer(1, 0.5 * cm))

        # Template phrases
        template_phrases = data.get('template_phrases', [])
        if template_phrases:
            story.append(Paragraph('Найденные шаблонные фразы', heading_style))
            for item in template_phrases:
                phrase = item.get('phrase', '')
                pos = item.get('pos', 0)
                story.append(Paragraph(f'• {phrase} (позиция: {pos})', normal_style))
            story.append(Spacer(1, 0.3 * cm))

        # Manipulation detector
        manipulation = data.get('manipulation', {})
        story.append(Paragraph('Детектор манипуляций', heading_style))
        zw = manipulation.get('zero_width', [])
        hs = manipulation.get('homoglyphs', [])
        us = manipulation.get('unusual_spaces', [])
        story.append(Paragraph(f'Нулевые пробелы: {len(zw)}', normal_style))
        story.append(Paragraph(f'Омоглифы: {len(hs)}', normal_style))
        story.append(Paragraph(f'Нестандартные пробелы: {len(us)}', normal_style))

        doc.build(story)

    def export_docx(self, filepath: str, data: dict):
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title = doc.add_heading('TextAnalyzer Pro — Отчёт', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Score section
        doc.add_heading('Метрики анализа', level=1)
        score = data.get('score', {})

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Shading Accent 1'
        hdr = table.rows[0].cells
        hdr[0].text = 'Метрика'
        hdr[1].text = 'Значение'

        rows = [
            ('Общий Score', str(score.get('total', '—'))),
            ('TTR', str(score.get('ttr', '—'))),
            ('Burstiness', str(score.get('burstiness', '—'))),
            ('Средняя длина предложения', str(score.get('avg_sentence_length', '—'))),
            ('Кол-во слов', str(score.get('word_count', '—'))),
            ('Уникальных слов', str(score.get('unique_words', '—'))),
            ('Кол-во предложений', str(score.get('sentence_count', '—'))),
        ]
        for label, value in rows:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value

        doc.add_paragraph()

        # Template phrases
        template_phrases = data.get('template_phrases', [])
        if template_phrases:
            doc.add_heading('Найденные шаблонные фразы', level=1)
            for item in template_phrases:
                phrase = item.get('phrase', '')
                pos = item.get('pos', 0)
                doc.add_paragraph(f'{phrase} (позиция: {pos})', style='List Bullet')

        # Manipulation detector
        doc.add_heading('Детектор манипуляций', level=1)
        manipulation = data.get('manipulation', {})
        zw = manipulation.get('zero_width', [])
        hs = manipulation.get('homoglyphs', [])
        us = manipulation.get('unusual_spaces', [])
        doc.add_paragraph(f'Нулевые пробелы: {len(zw)}')
        doc.add_paragraph(f'Омоглифы: {len(hs)}')
        doc.add_paragraph(f'Нестандартные пробелы: {len(us)}')

        doc.save(filepath)
